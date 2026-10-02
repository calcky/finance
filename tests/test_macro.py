"""Data-contract regressions: units, definitions, missingness and rolling history."""

import copy
from io import BytesIO
from pathlib import Path
from decimal import Decimal
import sys
import unittest

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from macro_common import number, observation
from macro_nbs import CONFIG, parse_json, parse_income
from macro_pbc import parse_workbook, parse_lpr, parse_repo, parse_fx, parse_yields
from macro_render import periods_between, load_topic, payload, document
from macro_catalog import TOPICS, SERIES
from sync_macro import merge, validate

URL = "https://example.org/source"


class SourceContracts(unittest.TestCase):
    def test_missing_not_zero_or_nonfinite(self):
        self.assertIsNone(number("—"))
        self.assertIsNone(observation("cpi_yoy", "2026-01", None, URL))
        self.assertEqual(observation("cpi_yoy", "2026-01", 0, URL)["value"], "0")
        for value in ("NaN", "Infinity", True):
            with self.assertRaises(ValueError):
                number(value)

    def test_dates_units_bounds(self):
        for key, period, value in [("cpi_yoy", "2026-13", 1), ("income_ytd", "2026-02", 1),
                                   ("repo_7d", "2026-02-30", 1), ("m1", "9999-01", 1),
                                   ("income_ytd", "9999-Q1", 1), ("pmi_manufacturing", "2026-01", 101)]:
            with self.assertRaises(ValueError):
                observation(key, period, value, URL)

    def nbs(self, key, value, period="202603MM", name=None, unit=None):
        _, ident, expected, expected_unit, transform, _ = CONFIG[key]
        if name is None:
            name = expected + (" (上年同月=100)" if transform == "index100" else "")
        return {"success": True, "data": [{"code": period, "values": [{"_id": ident, "da": "000000000000", "i_showname": name, "du_name": unit or expected_unit or "%", "value": value}]}]}

    def test_price_index_conversion_and_base(self):
        result = parse_json(self.nbs("cpi_yoy", "100.8"), ["cpi_yoy"], {})
        self.assertEqual(result[0]["value"], "0.8")
        with self.assertRaises(ValueError):
            parse_json(self.nbs("cpi_yoy", "100.8", name="居民消费价格指数 (上月=100)"), ["cpi_yoy"], {})
        self.assertEqual(parse_json(self.nbs("pmi_manufacturing", "49.8"), ["pmi_manufacturing"], {})[0]["value"], "49.8")

    def test_trade_unit_and_identity(self):
        result = parse_json(self.nbs("exports", "397851708"), ["exports"], {})
        self.assertEqual(result[0]["value"], "3978.51708")
        with self.assertRaises(ValueError):
            parse_json(self.nbs("exports", "397851708", unit="亿元"), ["exports"], {})
        data = self.nbs("exports", "10")
        data["data"][0]["values"][0]["da"] = "foreign"
        with self.assertRaises(ValueError):
            parse_json(data, ["exports"], {})

    def test_joint_january_february_not_single_month(self):
        data = self.nbs("industry_yoy", "5.0")
        data["data"] += self.nbs("industry_yoy", "5.1", "202602MM")["data"]
        result = parse_json(data, ["industry_yoy"], {})
        self.assertEqual([r["period"] for r in result], ["2026-03"])

    def test_duplicate_or_missing_indicator(self):
        data = self.nbs("cpi_yoy", "101")
        data["data"] *= 2
        with self.assertRaises(ValueError):
            parse_json(data, ["cpi_yoy"], {})
        with self.assertRaises(ValueError):
            parse_json({"success": True, "data": []}, ["cpi_yoy"], {})

    def test_income_is_ytd_and_real_is_explicit(self):
        html = '<meta name="PubDate" content="2026/07/15 10:00"><div class="txt-content"><table><tr><td>括号内实际增速</td></tr><tr><td>全国居民人均可支配收入</td><td>22981</td><td>5.2（4.2）</td></tr></table></div>'
        rows = parse_income(html, URL, "2026年上半年居民收入和消费支出情况")
        self.assertEqual({r["period"] for r in rows}, {"2026-Q2"})
        self.assertEqual(rows[-1]["value"], "4.2")
        self.assertEqual(rows[0]["published_at"], "2026-07-15")
        with self.assertRaises(ValueError):
            parse_income(html.replace("5.2（4.2）", "5.2"), URL, "2026年上半年居民收入和消费支出情况")

    def test_workbook_october_and_money_conversion(self):
        wb = Workbook()
        s = wb.active
        s.cell(3, 1, "单位：亿元")
        s.cell(8, 1, "M2")
        s.cell(10, 1, "M1")
        for month in range(1, 13):
            s.cell(6, month+3, float(f"2025.{month:02d}"))
            s.cell(8, month+3, 1000000+month)
            s.cell(10, month+3, 500000+month)
        stream = BytesIO()
        wb.save(stream)
        rows = parse_workbook(stream.getvalue(), "money", 2025, URL)
        october = next(r for r in rows if r["series_id"] == "m2" and r["period"] == "2025-10")
        self.assertEqual(october["value"], "100.001")
        with self.assertRaises(ValueError):
            parse_workbook(stream.getvalue(), "money", 2024, URL)

    def test_lpr_observation_not_publication(self):
        rows = parse_lpr({"data": {"columns": ["date", "open", "high", "low", "close", "volume", "1Y", "5Y"],
                                  "csv": "2025-05-30,,,,,,3.00,3.50\\r\\n"}}, URL)
        self.assertEqual(rows[0]["period"], "2025-05")
        self.assertEqual(rows[0]["published_at"], "")
        self.assertIn("2025-05-30", rows[0]["note"])

    def test_repo_tenor_and_no_operation(self):
        html = '2026-09-23<table><tr><td>期限</td><td>操作利率</td></tr><tr><td>14天</td><td>1.60%</td></tr><tr><td>7天</td><td>1. 40 %</td></tr></table>'
        rows = parse_repo(html, URL)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["value"], "1.40")
        self.assertEqual(parse_repo('2026-09-30<table><tr><td>期限</td><td>中标量</td></tr><tr><td>7天</td><td>0亿元</td></tr></table>', URL), [])

    def test_fx_joins_by_id_not_position(self):
        xml = '<chart><xaxis><value xid="a">30 Sep 2026</value><value xid="b">29 Sep 2026</value></xaxis><graphs><graph><value xid="b">6.74</value><value xid="a">6.7351</value></graph></graphs></chart>'
        rows = parse_fx(xml, URL)
        self.assertEqual({r["period"]: r["value"] for r in rows}, {"2026-09-29": "6.74", "2026-09-30": "6.7351"})

    def test_yield_selects_correct_tenors(self):
        html = '<table><tr><th>曲线</th><th>日期</th><th>10年</th><th>1年</th></tr><tr><td>中债国债收益率曲线</td><td>2026-09-30</td><td>1.6822</td><td>1.2197</td></tr></table>'
        rows = parse_yields(html, URL)
        self.assertEqual({r["series_id"]: r["value"] for r in rows}, {"yield_1y": "1.2197", "yield_10y": "1.6822"})

    def test_revision_rolling_history_and_withdrawal(self):
        old = [observation("repo_7d", f"2026-09-{day:02d}", "1.4", URL) for day in range(1, 4)]
        new = copy.deepcopy(old[1:])
        new[0]["value"] = "1.3"
        merged = merge(old, new, ["repo_7d"])
        self.assertEqual(len(merged), 3)
        self.assertEqual(merged[1]["value"], "1.3")
        with self.assertRaises(ValueError):
            merge(old, new[:1], ["repo_7d"])
        with self.assertRaises(ValueError):
            merge(old, [], ["repo_7d"])
        full = [observation("cpi_yoy", f"2026-{m:02d}", 1, URL) for m in (1, 2, 3)]
        with self.assertRaises(ValueError):
            merge(full, full[1:], ["cpi_yoy"])

    def test_calendar_gaps_and_quarters(self):
        self.assertEqual(periods_between("2025-12", "2026-02", "M"), ["2025-12", "2026-01", "2026-02"])
        self.assertEqual(periods_between("2025-Q4", "2026-Q2", "Q"), ["2025-Q4", "2026-Q1", "2026-Q2"])
        self.assertEqual(len(periods_between("2026-09-25", "2026-09-28", "D")), 4)


class CommittedSnapshot(unittest.TestCase):
    def test_economic_cross_checks(self):
        rows, _ = load_topic("trade-fx")
        values = {(r["series_id"], r["period"]): Decimal(r["value"]) for r in rows}
        for r in rows:
            if r["series_id"] == "trade_balance":
                p = r["period"]
                # Some historical official balances are rounded to 0.1/0.01
                # hundred-million USD; retain them, do not silently recompute.
                rounding = sum(Decimal(10) ** values[k, p].normalize().as_tuple().exponent / 2
                               for k in ("exports", "imports", "trade_balance"))
                self.assertLessEqual(abs(values["exports", p] - values["imports", p] - values["trade_balance", p]), rounding)
        rows, _ = load_topic("money-credit")
        values = {(r["series_id"], r["period"]): Decimal(r["value"]) for r in rows}
        for r in rows:
            if r["series_id"] == "m1":
                self.assertGreaterEqual(values["m2", r["period"]], Decimal(r["value"]))

    def test_all_topics_tables_and_payload_match(self):
        for slug, topic in TOPICS.items():
            rows, meta = load_topic(slug)
            validate(rows, meta["series"])
            data = payload(slug, rows, meta)
            text = document(slug, rows, meta)
            for chart in data["charts"]:
                self.assertEqual(len({SERIES[k]["frequency"] for k in chart["series"]}), 1)
                self.assertEqual(len({SERIES[k]["unit"] for k in chart["series"]}), 1)
                for key in chart["series"]:
                    actual = {r["period"]: float(r["value"]) for r in rows if r["series_id"] == key}
                    self.assertEqual(chart["values"][key], actual)
                self.assertIn(f'macro-{chart["id"]}-{chart["periods"][-1]}', text)


if __name__ == "__main__":
    unittest.main()
